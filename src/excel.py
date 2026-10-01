import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.utils import get_column_letter


# ============================================================
# PATHS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent
DATABASE_PATH = BASE_DIR / "data" / "library_analytics.db"
OUTPUT_PATH = BASE_DIR / "data" / "Library_Analytics.xlsx"


# ============================================================
# LOAD DATA FROM SQLITE
# ============================================================

def load_data():
    connection = sqlite3.connect(DATABASE_PATH)

    transactions = pd.read_sql_query(
        "SELECT * FROM borrow_transactions",
        connection
    )

    books = pd.read_sql_query(
        "SELECT * FROM books",
        connection
    )

    borrowers = pd.read_sql_query(
        "SELECT * FROM borrowers",
        connection
    )

    connection.close()

    return transactions, books, borrowers


# ============================================================
# PREPARE ANALYSIS TABLES
# ============================================================

def prepare_analysis(transactions, books, borrowers):

    tx = transactions.copy()

    # Convert date columns
    tx["borrow_date"] = pd.to_datetime(tx["borrow_date"])
    tx["due_date"] = pd.to_datetime(tx["due_date"])
    tx["return_date"] = pd.to_datetime(tx["return_date"])

    # -----------------------------
    # KPI SUMMARY
    # -----------------------------

    kpis = pd.DataFrame({
        "metric": [
            "Total Transactions",
            "Total Books",
            "Total Borrowers",
            "Returned Transactions",
            "Not Returned Transactions",
            "Total Fine",
            "Average Fine",
            "Maximum Fine"
        ],
        "value": [
            len(tx),
            len(books),
            len(borrowers),
            tx["return_date"].notna().sum(),
            tx["return_date"].isna().sum(),
            tx["fine"].sum(),
            tx["fine"].mean(),
            tx["fine"].max()
        ]
    })

    # -----------------------------
    # CATEGORY ANALYSIS
    # -----------------------------

    category = (
        tx.merge(
            books[["book_id", "category"]],
            on="book_id",
            how="left"
        )
        .groupby("category")
        .agg(
            borrow_count=("borrow_id", "size"),
            total_fine=("fine", "sum"),
            average_fine=("fine", "mean")
        )
        .reset_index()
        .sort_values("borrow_count", ascending=False)
    )

    # -----------------------------
    # MONTHLY TREND
    # -----------------------------

    monthly = (
        tx.assign(
            month=tx["borrow_date"].dt.to_period("M").astype(str)
        )
        .groupby("month")
        .size()
        .reset_index(name="borrow_count")
    )

    # -----------------------------
    # TOP 10 BOOKS
    # -----------------------------

    top_books = (
        tx.merge(
            books[["book_id", "title", "category"]],
            on="book_id",
            how="left"
        )
        .groupby(
            ["book_id", "title", "category"]
        )
        .size()
        .reset_index(name="borrow_count")
        .sort_values("borrow_count", ascending=False)
        .head(10)
    )

    # -----------------------------
    # TOP 10 BORROWERS
    # -----------------------------

    top_borrowers = (
        tx.merge(
            borrowers,
            on="borrower_id",
            how="left"
        )
        .groupby(
            ["borrower_id", "borrower_name"]
        )
        .agg(
            borrow_count=("borrow_id", "size"),
            total_fine=("fine", "sum")
        )
        .reset_index()
        .sort_values("borrow_count", ascending=False)
        .head(10)
    )

    # -----------------------------
    # RETURN STATUS
    # -----------------------------

    status = pd.DataFrame({
        "status": [
            "Returned",
            "Not Returned"
        ],
        "transaction_count": [
            tx["return_date"].notna().sum(),
            tx["return_date"].isna().sum()
        ]
    })

    status["percentage"] = (
        status["transaction_count"]
        / len(tx)
        * 100
    ).round(2)

    return (
        tx,
        kpis,
        category,
        monthly,
        top_books,
        top_borrowers,
        status
    )


# ============================================================
# CREATE EXCEL FILE
# ============================================================

def create_excel(
    transactions,
    books,
    borrowers,
    kpis,
    category,
    monthly,
    top_books,
    top_borrowers,
    status
):

    with pd.ExcelWriter(
        OUTPUT_PATH,
        engine="openpyxl"
    ) as writer:

        transactions.to_excel(
            writer,
            sheet_name="Transactions",
            index=False
        )

        books.to_excel(
            writer,
            sheet_name="Books",
            index=False
        )

        borrowers.to_excel(
            writer,
            sheet_name="Borrowers",
            index=False
        )

        kpis.to_excel(
            writer,
            sheet_name="KPIs",
            index=False
        )

        category.to_excel(
            writer,
            sheet_name="Category Analysis",
            index=False
        )

        monthly.to_excel(
            writer,
            sheet_name="Monthly Trend",
            index=False
        )

        top_books.to_excel(
            writer,
            sheet_name="Top Books",
            index=False
        )

        top_borrowers.to_excel(
            writer,
            sheet_name="Top Borrowers",
            index=False
        )

        status.to_excel(
            writer,
            sheet_name="Return Status",
            index=False
        )

        # Empty dashboard sheet
        pd.DataFrame().to_excel(
            writer,
            sheet_name="Dashboard",
            index=False
        )


# ============================================================
# FORMAT EXCEL
# ============================================================

def format_workbook():

    workbook = load_workbook(OUTPUT_PATH)

    # -----------------------------
    # General formatting
    # -----------------------------

    for worksheet in workbook.worksheets:

        worksheet.freeze_panes = "A2"

        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center"
            )

        # Automatic column widths
        for column in worksheet.columns:

            max_length = 0
            column_letter = get_column_letter(
                column[0].column
            )

            for cell in column:

                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max(max_length + 2, 10),
                30
            )

    # ========================================================
    # DASHBOARD
    # ========================================================

    dashboard = workbook["Dashboard"]

    dashboard["A1"] = "LIBRARY ANALYTICS DASHBOARD"

    dashboard["A1"].font = Font(
        bold=True,
        size=18
    )

    dashboard["A1"].alignment = Alignment(
        horizontal="center"
    )

    # -----------------------------
    # KPI section
    # -----------------------------

    kpi_sheet = workbook["KPIs"]

    dashboard["A3"] = "KEY METRICS"
    dashboard["A3"].font = Font(
        bold=True,
        size=14
    )

    for row in range(2, kpi_sheet.max_row + 1):

        metric = kpi_sheet.cell(
            row=row,
            column=1
        ).value

        value = kpi_sheet.cell(
            row=row,
            column=2
        ).value

        dashboard.cell(
            row=row + 3,
            column=1,
            value=metric
        )

        dashboard.cell(
            row=row + 3,
            column=2,
            value=value
        )

        dashboard.cell(
            row=row + 3,
            column=1
        ).font = Font(bold=True)

    # ========================================================
    # CHART 1 — BORROWINGS BY CATEGORY
    # ========================================================

    worksheet = workbook["Category Analysis"]

    chart = BarChart()

    chart.title = "Borrowings by Category"

    chart.y_axis.title = "Borrowings"

    chart.x_axis.title = "Category"

    data = Reference(
        worksheet,
        min_col=2,
        min_row=1,
        max_row=worksheet.max_row
    )

    categories = Reference(
        worksheet,
        min_col=1,
        min_row=2,
        max_row=worksheet.max_row
    )

    chart.add_data(
        data,
        titles_from_data=True
    )

    chart.set_categories(categories)

    chart.height = 8
    chart.width = 14

    dashboard.add_chart(
        chart,
        "D3"
    )

    # ========================================================
    # CHART 2 — MONTHLY BORROWING TREND
    # ========================================================

    worksheet = workbook["Monthly Trend"]

    chart = LineChart()

    chart.title = "Monthly Borrowing Trend"

    chart.y_axis.title = "Borrowings"

    chart.x_axis.title = "Month"

    data = Reference(
        worksheet,
        min_col=2,
        min_row=1,
        max_row=worksheet.max_row
    )

    categories = Reference(
        worksheet,
        min_col=1,
        min_row=2,
        max_row=worksheet.max_row
    )

    chart.add_data(
        data,
        titles_from_data=True
    )

    chart.set_categories(categories)

    chart.height = 8
    chart.width = 14

    dashboard.add_chart(
        chart,
        "D20"
    )

    # ========================================================
    # CHART 3 — RETURN STATUS
    # ========================================================

    worksheet = workbook["Return Status"]

    chart = PieChart()

    chart.title = "Return Status"

    data = Reference(
        worksheet,
        min_col=2,
        min_row=1,
        max_row=worksheet.max_row
    )

    categories = Reference(
        worksheet,
        min_col=1,
        min_row=2,
        max_row=worksheet.max_row
    )

    chart.add_data(
        data,
        titles_from_data=True
    )

    chart.set_categories(categories)

    chart.height = 8
    chart.width = 12

    dashboard.add_chart(
        chart,
        "T3"
    )

    # ========================================================
    # CHART 4 — TOP BOOKS
    # ========================================================

    worksheet = workbook["Top Books"]

    chart = BarChart()

    chart.title = "Top 10 Most Borrowed Books"

    chart.y_axis.title = "Borrowings"

    chart.x_axis.title = "Book"

    data = Reference(
        worksheet,
        min_col=4,
        min_row=1,
        max_row=worksheet.max_row
    )

    categories = Reference(
        worksheet,
        min_col=2,
        min_row=2,
        max_row=worksheet.max_row
    )

    chart.add_data(
        data,
        titles_from_data=True
    )

    chart.set_categories(categories)

    chart.height = 10
    chart.width = 16

    dashboard.add_chart(
        chart,
        "T20"
    )

    # ========================================================
    # SAVE
    # ========================================================

    workbook.save(OUTPUT_PATH)


# ============================================================
# MAIN
# ============================================================

def main():

    print("Loading library database...")

    transactions, books, borrowers = load_data()

    print("Database loaded.")

    (
        tx,
        kpis,
        category,
        monthly,
        top_books,
        top_borrowers,
        status
    ) = prepare_analysis(
        transactions,
        books,
        borrowers
    )

    print("Analysis tables prepared.")

    create_excel(
        tx,
        books,
        borrowers,
        kpis,
        category,
        monthly,
        top_books,
        top_borrowers,
        status
    )

    format_workbook()

    print()
    print("===================================")
    print("Excel workbook created successfully")
    print("===================================")
    print(f"File: {OUTPUT_PATH}")
    print()
    print("Sheets created:")
    print("1. Transactions")
    print("2. Books")
    print("3. Borrowers")
    print("4. KPIs")
    print("5. Category Analysis")
    print("6. Monthly Trend")
    print("7. Top Books")
    print("8. Top Borrowers")
    print("9. Return Status")
    print("10. Dashboard")


if __name__ == "__main__":
    main()